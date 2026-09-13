"""
Excel Formula Recalculation using LibreOffice.

Recalculates all formulas in an Excel file and scans for errors.
Runs on the AgentCore runtime container where LibreOffice is available.
"""

import logging
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

logger = logging.getLogger(__name__)

EXCEL_ERRORS = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]


def _scan_errors(filename: str) -> dict:
    """Scan recalculated file for Excel formula errors and count formulas."""
    from openpyxl import load_workbook

    error_details = {err: [] for err in EXCEL_ERRORS}
    total_errors = 0
    formula_count = 0

    try:
        wb = load_workbook(filename, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and isinstance(cell.value, str):
                        for err in EXCEL_ERRORS:
                            if err in cell.value:
                                error_details[err].append(f"{sheet_name}!{cell.coordinate}")
                                total_errors += 1
                                break
        wb.close()

        wb_formulas = load_workbook(filename, data_only=False)
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_count += 1
        wb_formulas.close()

    except Exception as e:
        logger.error(f"Error scanning spreadsheet: {e}")
        return {"status": "scan_error", "error": str(e)}

    # Empty numeric formula caches are not proof of recalculation. Excel can
    # calculate these on opening, while previews/readers may display blanks.
    missing_cached_values = 0
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(filename) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                continue
            for cell in ET.fromstring(archive.read(name)).findall(".//x:c", ns):
                if cell.find("x:f", ns) is None:
                    continue
                value = cell.find("x:v", ns)
                if value is None or (value.text is None and cell.get("t") != "str"):
                    missing_cached_values += 1

    result = {
        "status": "incomplete" if missing_cached_values else ("success" if total_errors == 0 else "errors_found"),
        "missing_cached_values": missing_cached_values,
        "total_errors": total_errors,
        "total_formulas": formula_count,
    }

    if total_errors > 0:
        result["error_summary"] = {}
        for err_type, locations in error_details.items():
            if locations:
                result["error_summary"][err_type] = {
                    "count": len(locations),
                    "locations": locations[:10],
                }

    return result


def recalc_spreadsheet(
    file_bytes: bytes,
    filename: str = "temp.xlsx",
    timeout: int = 30,
) -> tuple[bytes, dict]:
    """Open and save with Calc in an isolated profile, then verify formula caches."""
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / Path(filename).name
        source.write_bytes(file_bytes)
        output_dir = root / "recalculated"
        output_dir.mkdir()
        command = [
            "soffice",
            f"-env:UserInstallation={(root / 'profile').as_uri()}",
            "--headless", "--norestore",
            "--convert-to", "xlsx:Calc MS Excel 2007 XML",
            "--outdir", str(output_dir), str(source),
        ]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            logger.warning("LibreOffice recalculation timed out")
            return file_bytes, {"status": "skipped", "reason": "timeout"}
        except FileNotFoundError:
            return file_bytes, {"status": "skipped", "reason": "soffice_not_found"}

        output = output_dir / source.name
        if result.returncode != 0 or not output.is_file():
            logger.warning("LibreOffice recalculation failed: %s", result.stderr[:200])
            return file_bytes, {"status": "skipped", "reason": "recalc_failed"}
        report = _scan_errors(str(output))
        if report["status"] not in {"success", "errors_found"}:
            return file_bytes, report
        return output.read_bytes(), report
